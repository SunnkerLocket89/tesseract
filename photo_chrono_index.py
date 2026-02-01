import datetime
import os
import re

import cv2
import fitz
import numpy as np
import pandas as pd
import pytesseract
from openpyxl import Workbook

PDFS = {
    "Batch1": r"OCR_D22000030 PHOTOS.pdf",
    "Batch2": r"OCR_D22000030_2_PHOTOS_BATCH_(RELEASED 08-29-25).pdf",
    "Batch3": r"OCR_D22000030_3_PHOTOS_BATCH_(RELEASED 09-22-25).pdf",
    "Batch4": r"3_COMPRESSED_D22000030 PHOTOS BATCH 4 (RELEASED 01-20-26).pdf",
}

SLATES_CSV = r"_slate_candidates.csv"
OUT_XLSX = r"photo_chrono_index_NEAR_DEFINITIVE.xlsx"

date_pat = re.compile(r"(\d{1,2})[\/\-](\d{1,2})[\/\-](\d{2,4})")
time_pat = re.compile(r"(\d{1,2})\s*:\s*(\d{2})(?:\s*:\s*(\d{2}))?\s*(AM|PM|am|pm)?")


def parse_dt(txt: str):
    txt = txt or ""
    d = None
    m = date_pat.search(txt)
    if m:
        mo, da, yy = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if yy < 100:
            yy += 2000 if yy <= 60 else 1900
        try:
            d = datetime.date(yy, mo, da)
        except ValueError:
            d = None

    t = None
    m = time_pat.search(txt)
    if m:
        hh, mi = int(m.group(1)), int(m.group(2))
        ss = int(m.group(3)) if m.group(3) else 0
        ampm = (m.group(4) or "").lower()
        if ampm == "pm" and hh < 12:
            hh += 12
        if ampm == "am" and hh == 12:
            hh = 0
        try:
            t = datetime.time(hh, mi, ss)
        except ValueError:
            t = None

    dt = datetime.datetime.combine(d, t) if (d and t) else None
    return dt, d, t


def render_center_gray(page, dpi=150):
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    if img.shape[2] == 4:
        img = img[:, :, :3]
    gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
    h, w = gray.shape
    y0, y1 = int(h * 0.05), int(h * 0.95)
    x0, x1 = int(w * 0.05), int(w * 0.95)
    return gray[y0:y1, x0:x1]


def ocr_fast(gray):
    if gray.shape[0] < 1000:
        s = 1000 / gray.shape[0]
        gray = cv2.resize(
            gray,
            (int(gray.shape[1] * s), int(gray.shape[0] * s)),
            interpolation=cv2.INTER_CUBIC,
        )
    gray = cv2.equalizeHist(gray)
    _, th = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    try:
        txt = pytesseract.image_to_string(th, config="--psm 6", timeout=3)
    except RuntimeError:
        txt = ""
    return re.sub(r"\s+", " ", txt).strip()


def main():
    slates = pd.read_csv(SLATES_CSV)
    slates = slates.sort_values(["batch", "page_number"]).reset_index(drop=True)

    docs = {}
    for b, fn in PDFS.items():
        if not os.path.exists(fn):
            raise FileNotFoundError(f"Missing PDF: {fn}")
        docs[b] = fitz.open(fn)

    chrono_rows = []
    global_idx = 0
    for b, doc in docs.items():
        for p in range(doc.page_count):
            global_idx += 1
            chrono_rows.append(
                {
                    "batch": b,
                    "pdf_file": os.path.basename(PDFS[b]),
                    "page_number": p + 1,
                    "global_index": global_idx,
                }
            )
    chrono = pd.DataFrame(chrono_rows)

    slate_key = set(zip(slates["batch"], slates["page_number"].astype(int)))
    chrono["is_slate_candidate"] = chrono.apply(
        lambda r: (r["batch"], r["page_number"]) in slate_key,
        axis=1,
    )

    ocr_out = []
    for _, r in slates.iterrows():
        b = r["batch"]
        doc = docs[b]
        page = doc.load_page(int(r["page_number"]) - 1)

        text_layer = re.sub(r"\s+", " ", (page.get_text("text") or "")).strip()
        if b != "Batch4" and len(text_layer) > 30:
            best_text = text_layer
            source = "text_layer"
        else:
            gray = render_center_gray(page, dpi=150)
            best_text = ocr_fast(gray)
            source = "image_ocr"

        dt, d, t = parse_dt(best_text)
        ocr_out.append(
            {
                "batch": b,
                "page_number": int(r["page_number"]),
                "slate_score": float(r.get("slate_score", 0)),
                "best_source": source,
                "best_text": best_text[:2000],
                "auto_datetime": dt,
                "auto_date": d,
                "auto_time": t,
            }
        )
    slate_ocr = pd.DataFrame(ocr_out)

    chrono["run_id"] = 0
    for b in chrono["batch"].unique():
        sub = chrono[chrono["batch"] == b].sort_values("page_number").copy()
        run = 0
        run_ids = []
        for _, rr in sub.iterrows():
            if rr["is_slate_candidate"]:
                run += 1
            run_ids.append(run if run else 1)
        chrono.loc[sub.index, "run_id"] = run_ids

    chrono["index_in_run"] = chrono.groupby(["batch", "run_id"]).cumcount()

    chrono = chrono.merge(
        slate_ocr[
            [
                "batch",
                "page_number",
                "slate_score",
                "best_source",
                "best_text",
                "auto_datetime",
                "auto_date",
                "auto_time",
            ]
        ],
        on=["batch", "page_number"],
        how="left",
    )

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Chronology_Working"
    for c, col in enumerate(chrono.columns, start=1):
        ws1.cell(row=1, column=c, value=col)
    for r_i, row in enumerate(chrono.itertuples(index=False), start=2):
        for c_i, val in enumerate(row, start=1):
            ws1.cell(row=r_i, column=c_i, value=val)

    ws2 = wb.create_sheet("Slate_Anchors")
    anchors = chrono[chrono["is_slate_candidate"]].sort_values(
        ["batch", "run_id", "page_number"]
    )
    anchors = anchors.groupby(["batch", "run_id"]).first().reset_index()
    anchors["Manual_AnchorDateTime"] = ""
    anchor_cols = [
        "batch",
        "run_id",
        "pdf_file",
        "page_number",
        "slate_score",
        "best_source",
        "auto_datetime",
        "auto_date",
        "auto_time",
        "best_text",
        "Manual_AnchorDateTime",
    ]
    ws2.append(anchor_cols)
    for _, r in anchors[anchor_cols].iterrows():
        ws2.append(list(r.values))

    ws3 = wb.create_sheet("Run_Review_Pack")
    ws3.append(
        [
            "batch",
            "run_id",
            "page_number",
            "slate_score",
            "best_source",
            "auto_datetime",
            "auto_date",
            "best_text",
        ]
    )
    for _, r in anchors.iterrows():
        ws3.append(
            [
                r["batch"],
                r["run_id"],
                r["page_number"],
                r.get("slate_score", ""),
                r.get("best_source", ""),
                r.get("auto_datetime", ""),
                r.get("auto_date", ""),
                (r.get("best_text", "") or "")[:300],
            ]
        )

    wb.save(OUT_XLSX)

    for d in docs.values():
        d.close()

    print("Wrote:", OUT_XLSX)


if __name__ == "__main__":
    main()
